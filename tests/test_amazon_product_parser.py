from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict

import pytest
from openpyxl import Workbook, load_workbook

from src.amazon_product_parser import (
    AmazonProductParser,
    PlaywrightAmazonProductParser,
    create_parser,
    process_excel,
)

FIXTURES = Path(__file__).parent / "data"


def load_fixture(name: str) -> str:
    return (FIXTURES / name).read_text(encoding="utf-8")


def test_create_parser_static_returns_default() -> None:
    parser = create_parser("static")
    assert isinstance(parser, AmazonProductParser)
    assert not isinstance(parser, PlaywrightAmazonProductParser)


def test_create_parser_playwright_returns_subclass() -> None:
    parser = create_parser("playwright")
    assert isinstance(parser, PlaywrightAmazonProductParser)


def test_playwright_parser_requires_dependency(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("src.amazon_product_parser.sync_playwright", None, raising=False)
    parser = PlaywrightAmazonProductParser(fallback_to_static=False)
    with pytest.raises(RuntimeError):
        parser.parse("https://www.amazon.com/dp/B012345678")


def test_playwright_parser_falls_back_to_static(monkeypatch: pytest.MonkeyPatch) -> None:
    html = load_fixture("sample_product.html")

    def fail_render(self: PlaywrightAmazonProductParser, url: str) -> str:
        raise RuntimeError("render failed")

    def fake_static_parse(self: AmazonProductParser, url: str) -> Any:  # type: ignore[override]
        return AmazonProductParser().parse_html(html, url=url)

    monkeypatch.setattr(
        "src.amazon_product_parser.PlaywrightAmazonProductParser._render_url",
        fail_render,
        raising=False,
    )
    monkeypatch.setattr(
        "src.amazon_product_parser.AmazonProductParser.parse",
        fake_static_parse,
        raising=False,
    )

    parser = PlaywrightAmazonProductParser()
    product = parser.parse("https://www.amazon.com/dp/B012345678")
    assert product.title == "Sample Product Title"


def test_parse_full_product_html() -> None:
    html = load_fixture("sample_product.html")
    parser = AmazonProductParser()

    product = parser.parse_html(html, url="https://www.amazon.com/dp/B012345678")

    assert product.title == "Sample Product Title"
    assert product.asin == "B012345678"
    assert product.price == pytest.approx(19.99)
    assert product.currency == "USD"
    assert product.rating == pytest.approx(4.5)
    assert product.review_count == 1234
    assert product.image == "https://images.example.com/sample-hero-large.jpg"

    assert product.bullet_points == [
        "Bullet one feature",
        "Bullet two feature",
        "Bullet three feature",
    ]

    assert product.details["Brand"] == "SampleBrand"
    assert product.details["ASIN"] == "B012345678"
    assert product.details["Item Weight"] == "1.2 pounds"
    assert product.details["Material"] == "Plastic"

    as_dict = product.to_dict()
    assert json.loads(json.dumps(as_dict, ensure_ascii=False)) == as_dict


def test_parse_minimal_product_html() -> None:
    html = load_fixture("sample_product_minimal.html")
    parser = AmazonProductParser()

    product = parser.parse_html(html, url="https://www.amazon.co.uk/dp/B0FEDCBA98")

    assert product.title == "Minimal Product"
    assert product.asin == "B0FEDCBA98"
    assert product.price == pytest.approx(10.49)
    assert product.currency == "GBP"
    assert product.rating == pytest.approx(3.7)
    assert product.review_count == 87
    assert product.bullet_points == []
    assert product.details["ASIN"] == "B0FEDCBA98"


def test_parse_uses_custom_session_headers() -> None:
    class FakeResponse:
        def __init__(self, text: str) -> None:
            self.text = text

        def raise_for_status(self) -> None:  # pragma: no cover - simple stub
            return None

    class FakeSession:
        def __init__(self, text: str) -> None:
            self._text = text
            self.headers: Dict[str, str] = {}
            self.captured: Dict[str, Any] = {}

        def get(self, url: str, *, headers: Dict[str, str], timeout: float) -> FakeResponse:
            self.captured = {"url": url, "headers": headers, "timeout": timeout}
            return FakeResponse(self._text)

    session = FakeSession(load_fixture("sample_product_minimal.html"))
    parser = AmazonProductParser(session=session, timeout=9.5)

    product = parser.parse("https://www.amazon.co.uk/dp/B0FEDCBA98")

    assert product.title == "Minimal Product"
    assert session.captured["url"] == "https://www.amazon.co.uk/dp/B0FEDCBA98"
    # Ensure the parser always sends a deterministic user agent header
    assert "User-Agent" in session.captured["headers"]
    assert session.captured["timeout"] == 9.5


def test_process_excel_updates_workbook(tmp_path: Path) -> None:
    workbook_path = tmp_path / "input.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet["A1"] = "Amazon URL"
    sheet["A2"] = "https://www.amazon.com/dp/B012345678"
    sheet["A3"] = "https://www.amazon.co.uk/dp/B0FEDCBA98"
    workbook.save(workbook_path)
    workbook.close()

    parser = AmazonProductParser()

    class FixtureParser:
        def parse(self, url: str):  # type: ignore[override]
            if "B012345678" in url:
                return parser.parse_html(load_fixture("sample_product.html"), url=url)
            return parser.parse_html(load_fixture("sample_product_minimal.html"), url=url)

    destination = process_excel(workbook_path, parser=FixtureParser())
    assert Path(destination) == workbook_path

    result_workbook = load_workbook(workbook_path)
    result_sheet = result_workbook.active

    assert result_sheet["B1"].value == "ASIN"
    assert result_sheet["C1"].value == "Title"

    assert result_sheet["B2"].value == "B012345678"
    assert result_sheet["C2"].value == "Sample Product Title"
    assert result_sheet["D2"].value == pytest.approx(19.99)
    assert result_sheet["E2"].value == "USD"
    assert result_sheet["F2"].value == pytest.approx(4.5)
    assert result_sheet["G2"].value == 1234
    assert result_sheet["I2"].value.startswith("Bullet one feature")
    assert json.loads(result_sheet["J2"].value)["Brand"] == "SampleBrand"
    assert result_sheet["K2"].value is None

    assert result_sheet["B3"].value == "B0FEDCBA98"
    assert result_sheet["D3"].value == pytest.approx(10.49)
    assert result_sheet["E3"].value == "GBP"
    assert result_sheet["F3"].value == pytest.approx(3.7)
    assert result_sheet["G3"].value == 87
    assert result_sheet["I3"].value is None
    assert result_sheet["K3"].value is None

    result_workbook.close()


def test_process_excel_accepts_asin_values(tmp_path: Path) -> None:
    workbook_path = tmp_path / "input_asin.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet["A1"] = "ASIN"
    sheet["A2"] = "B012345678"
    sheet["A3"] = "B0FEDCBA98"
    workbook.save(workbook_path)
    workbook.close()

    parser = AmazonProductParser()
    captured_urls: list[str] = []

    class FixtureParser:
        def parse(self, url: str):  # type: ignore[override]
            captured_urls.append(url)
            if "B012345678" in url:
                return parser.parse_html(load_fixture("sample_product.html"), url=url)
            return parser.parse_html(load_fixture("sample_product_minimal.html"), url=url)

    destination = process_excel(workbook_path, parser=FixtureParser())
    assert Path(destination) == workbook_path

    assert captured_urls == [
        "https://www.amazon.com/dp/B012345678",
        "https://www.amazon.com/dp/B0FEDCBA98",
    ]

    result_workbook = load_workbook(workbook_path)
    result_sheet = result_workbook.active

    assert result_sheet["B1"].value == "ASIN"
    assert result_sheet["C1"].value == "Title"
    assert result_sheet["B2"].value == "B012345678"
    assert result_sheet["C2"].value == "Sample Product Title"
    assert result_sheet["B3"].value == "B0FEDCBA98"
    assert result_sheet["C3"].value == "Minimal Product"

    result_workbook.close()
